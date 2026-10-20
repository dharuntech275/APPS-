"""
build_presentation.py
Generates a 10-slide, professional widescreen (16:9) PowerPoint presentation (.pptx)
for the Final Year Mini Project:
"AI PREDICTIVE PERSONNEL STRESS AND WELFARE MONITORING SYSTEMS"
Presented by: DHARUN.M (Final Year, Computer Science and Business Systems)

Constraints strictly met:
- Exactly 10 slides.
- Slide 1 has:
    NAME : DHARUN.M
    YEAR : FINAL YEAR
    DEPT : COMPUTER SCIENCE AND BUSINESS SYSTEMS
    PROJECT TITLE : AI PREDICTIVE PERSONNEL STRESS AND WELFARE MONITORING SYSTEMS
- Slides 2 to 10 have strictly 3 clear, high-impact bullet points each.
- Professional English phrasing, clean defense-grade visual styling.
"""

import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation(output_path="APPS_Mini_Project_Presentation.pptx"):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette Definitions
    c_navy_dark = RGBColor(15, 23, 42)      # #0F172A - Deep Navy
    c_navy_card = RGBColor(30, 41, 59)      # #1E293B - Dark Slate Card
    c_royal_blue = RGBColor(37, 99, 235)    # #2563EB - Royal Blue Accent
    c_teal = RGBColor(13, 148, 136)         # #0D9488 - Deep Teal
    c_white = RGBColor(255, 255, 255)       # White
    c_slate_light = RGBColor(241, 245, 249) # #F1F5F9 - Card Background
    c_slate_border = RGBColor(203, 213, 225)# #CBD5E1 - Border
    c_slate_text = RGBColor(51, 65, 85)     # #334155 - Body Text
    c_muted = RGBColor(100, 116, 139)       # #64748B - Muted Gray
    c_crimson = RGBColor(220, 38, 38)       # #DC2626 - Alert Crimson
    c_bg_slide = RGBColor(248, 250, 252)    # #F8FAFC - Slide BG

    blank_layout = prs.slide_layouts[6]

    # Helper: add background rectangle
    def set_bg(slide, color):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    # =========================================================================
    # SLIDE 1: TITLE SLIDE (Dark Executive Theme)
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_bg(s1, c_navy_dark)

    # Top Pill Tag
    tag_box = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(0.8), Inches(4.2), Inches(0.45))
    tag_box.fill.solid()
    tag_box.fill.fore_color.rgb = RGBColor(30, 58, 138)
    tag_box.line.color.rgb = c_royal_blue
    tag_tf = tag_box.text_frame
    tag_tf.text = "FINAL YEAR MINI PROJECT PRESENTATION"
    p = tag_tf.paragraphs[0]
    p.font.name = "Calibri"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = RGBColor(147, 197, 253)
    p.alignment = PP_ALIGN.CENTER

    # Project Title
    t_box = s1.shapes.add_textbox(Inches(1.0), Inches(1.4), Inches(11.333), Inches(1.8))
    t_tf = t_box.text_frame
    t_tf.word_wrap = True
    p_title = t_tf.paragraphs[0]
    p_title.text = "AI PREDICTIVE PERSONNEL STRESS AND WELFARE MONITORING SYSTEMS"
    p_title.font.name = "Calibri"
    p_title.font.size = Pt(28)
    p_title.font.bold = True
    p_title.font.color.rgb = c_white
    p_title.alignment = PP_ALIGN.LEFT

    # Subtitle line
    p_sub = t_tf.add_paragraph()
    p_sub.text = "A Behavioral Telemetry & Clinical AI Early-Warning Triage Platform for Armed Forces"
    p_sub.font.name = "Calibri"
    p_sub.font.size = Pt(14)
    p_sub.font.color.rgb = RGBColor(148, 163, 184)
    p_sub.space_before = Pt(8)

    # Student Details Card
    card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(3.6), Inches(11.333), Inches(2.9))
    card.fill.solid()
    card.fill.fore_color.rgb = c_navy_card
    card.line.color.rgb = c_royal_blue
    card.line.width = Pt(1.5)

    c_tf = card.text_frame
    c_tf.word_wrap = True
    c_tf.vertical_anchor = MSO_ANCHOR.TOP

    # Header inside card
    p_card_hd = c_tf.paragraphs[0]
    p_card_hd.text = "CANDIDATE & ACADEMIC CREDENTIALS"
    p_card_hd.font.name = "Calibri"
    p_card_hd.font.size = Pt(12)
    p_card_hd.font.bold = True
    p_card_hd.font.color.rgb = RGBColor(56, 189, 248)
    p_card_hd.space_after = Pt(14)

    # Mandatory details
    details = [
        ("NAME :", "DHARUN.M"),
        ("YEAR :", "FINAL YEAR"),
        ("DEPT :", "COMPUTER SCIENCE AND BUSINESS SYSTEMS"),
        ("PROJECT TITLE :", "AI PREDICTIVE PERSONNEL STRESS AND WELFARE MONITORING SYSTEMS"),
    ]

    for label, val in details:
        p_row = c_tf.add_paragraph()
        p_row.text = f"{label:<16}  {val}"
        p_row.font.name = "Consolas"
        p_row.font.size = Pt(13)
        p_row.font.bold = True
        p_row.font.color.rgb = c_white
        p_row.space_after = Pt(6)

    # Helper function for Content Slides (Slides 2 to 10)
    # Strictly 3 cards/points per slide
    def add_content_slide(slide_num, category, title, points):
        s = prs.slides.add_slide(blank_layout)
        set_bg(s, c_bg_slide)

        # Header Pill
        cat_box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(0.5), Inches(3.2), Inches(0.35))
        cat_box.fill.solid()
        cat_box.fill.fore_color.rgb = RGBColor(239, 246, 255)
        cat_box.line.color.rgb = RGBColor(191, 219, 254)
        c_p = cat_box.text_frame.paragraphs[0]
        c_p.text = f"SLIDE {slide_num:02d} | {category.upper()}"
        c_p.font.name = "Calibri"
        c_p.font.size = Pt(9.5)
        c_p.font.bold = True
        c_p.font.color.rgb = c_royal_blue
        c_p.alignment = PP_ALIGN.CENTER

        # Slide Title
        t_box = s.shapes.add_textbox(Inches(0.9), Inches(0.9), Inches(11.5), Inches(0.8))
        t_tf = t_box.text_frame
        t_p = t_tf.paragraphs[0]
        t_p.text = title
        t_p.font.name = "Calibri"
        t_p.font.size = Pt(22)
        t_p.font.bold = True
        t_p.font.color.rgb = c_navy_dark

        # 3 Structured Horizontal Cards
        card_w = Inches(11.533)
        card_h = Inches(1.55)
        start_y = Inches(1.85)
        gap_y = Inches(0.2)

        for i, (p_title, p_desc) in enumerate(points):
            y_pos = start_y + (i * (card_h + gap_y))

            # Outer Card
            c_box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), y_pos, card_w, card_h)
            c_box.fill.solid()
            c_box.fill.fore_color.rgb = c_white
            c_box.line.color.rgb = c_slate_border
            c_box.line.width = Pt(1)

            # Left Accent Badge (01, 02, 03)
            b_box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.15), y_pos + Inches(0.25), Inches(0.9), Inches(0.95))
            b_box.fill.solid()
            # Alternating colors: Royal Blue, Teal, Navy
            accent_cols = [c_royal_blue, c_teal, c_navy_dark]
            b_box.fill.fore_color.rgb = accent_cols[i % 3]
            b_box.line.fill.background()

            b_tf = b_box.text_frame
            b_p = b_tf.paragraphs[0]
            b_p.text = f"{i+1:02d}"
            b_p.font.name = "Calibri"
            b_p.font.size = Pt(20)
            b_p.font.bold = True
            b_p.font.color.rgb = c_white
            b_p.alignment = PP_ALIGN.CENTER

            # Text box for Card Title & Content
            txt_box = s.shapes.add_textbox(Inches(2.25), y_pos + Inches(0.15), Inches(10.0), Inches(1.25))
            txt_tf = txt_box.text_frame
            txt_tf.word_wrap = True

            pt_head = txt_tf.paragraphs[0]
            pt_head.text = p_title
            pt_head.font.name = "Calibri"
            pt_head.font.size = Pt(14)
            pt_head.font.bold = True
            pt_head.font.color.rgb = c_navy_dark

            pt_body = txt_tf.add_paragraph()
            pt_body.text = p_desc
            pt_body.font.name = "Calibri"
            pt_body.font.size = Pt(11.5)
            pt_body.font.color.rgb = c_slate_text
            pt_body.space_before = Pt(4)

        # Footer
        f_box = s.shapes.add_textbox(Inches(0.9), Inches(7.0), Inches(11.5), Inches(0.35))
        f_tf = f_box.text_frame
        f_p = f_tf.paragraphs[0]
        f_p.text = "APPS - AI Predictive Personnel Stress & Welfare Monitoring Systems | Computer Science & Business Systems"
        f_p.font.name = "Calibri"
        f_p.font.size = Pt(9)
        f_p.font.color.rgb = c_muted

        return s

    # =========================================================================
    # SLIDE 2: PROBLEM STATEMENT
    # =========================================================================
    add_content_slide(
        2,
        "Problem Statement",
        "Challenges in Armed Forces Psychological Welfare",
        [
            (
                "Undetected Combat Fatigue & Operational Strain",
                "Armed forces personnel face prolonged high-intensity combat exposure and erratic sleep schedules, leading to cumulative mental exhaustion that frequently goes unnoticed."
            ),
            (
                "Infrequent Manual Evaluations & Fatal Delays",
                "Traditional military welfare relies on sporadic, in-person psychological reviews that create critical response delays in suicide prevention and mental trauma intervention."
            ),
            (
                "Absence of Confidential Real-Time Triage Systems",
                "Soldiers often hesitate to report mental health struggles due to social stigma, while commanding officers lack automated tools to detect high-risk personnel early."
            )
        ]
    )

    # =========================================================================
    # SLIDE 3: PROJECT OBJECTIVES
    # =========================================================================
    add_content_slide(
        3,
        "Objectives",
        "Core Project Goals & Clinical Engineering Vision",
        [
            (
                "Proactive AI-Driven Early Stress Detection",
                "Engineer an automated behavioral telemetry and clinical monitoring system capable of predicting soldier psychological burnout before it leads to critical operational failure."
            ),
            (
                "Automated Multi-Metric Triage Categorization",
                "Implement algorithmic classification dividing personnel into High, Medium, and Low risk tiers based on daily sleep, duty stress, heart rate, and squad morale."
            ),
            (
                "Instant Frontline De-Escalation & Crisis Dispatch",
                "Provide immediate psychological stabilization through interactive 4-7-8 breathing exercises and direct 1-click Crisis SOS routing to military helpline 1902 and Tele-MANAS."
            )
        ]
    )

    # =========================================================================
    # SLIDE 4: SYSTEM ARCHITECTURE
    # =========================================================================
    add_content_slide(
        4,
        "Architecture",
        "3-Tier Role Hierarchy & Cloud Infrastructure",
        [
            (
                "Tri-Level Military Role-Based Access Control (RBAC)",
                "Strict operational separation across Personnel (Daily survey & SOS), Welfare Officers (Triage roster & Clinical notes), and Senior Command HQ (Analytics & PIN control)."
            ),
            (
                "Serverless Backend with Vercel Cloud Blob Persistence",
                "Python Flask WSGI microservice integrated with private Vercel Blob storage, ensuring zero database loss across ephemeral cloud cold-boots through automated snapshot hydration."
            ),
            (
                "Dual-Screen Responsive Design (Laptop & Mobile)",
                "Executive 2-column laptop dashboard (1.15fr : 0.85fr grid) paired with a touch-optimized mobile slide-out drawer navigation with zero horizontal overflow."
            )
        ]
    )

    # =========================================================================
    # SLIDE 5: PERSONNEL MODULE
    # =========================================================================
    add_content_slide(
        5,
        "Personnel Module",
        "Soldier Daily Self-Assessment & Biometric Tracking",
        [
            (
                "4-Dimensional Daily Health Telemetry Capture",
                "Personnel log daily sleep duration (0-12h), operational duty stress (1-10 dynamic slider), physical resting heart rate (bpm), and squad/family support level."
            ),
            (
                "Smartwatch & BLE Sensor Simulation Engine",
                "One-click telemetry synchronization that automatically loads realistic wearable sensor data into daily assessment fields, reducing soldier reporting friction."
            ),
            (
                "Gamified Morale Badges & Habit Formation Streaks",
                "Unlocks military achievement milestones such as 'Mental Resilience Sentinel' and tracks consecutive daily streaks to foster positive mental self-care habits."
            )
        ]
    )

    # =========================================================================
    # SLIDE 6: CLINICAL TRIAGE & FATIGUE FORECASTER
    # =========================================================================
    add_content_slide(
        6,
        "Clinical Intelligence",
        "Automated Risk Triage & Sleep Debt Analytics",
        [
            (
                "Rule-Based Clinical Triage Scoring Algorithm",
                "Instantly evaluates inputs against medical thresholds to tag entries as High Risk (Red - immediate intervention), Medium Risk (Yellow), or Low Risk (Green - stable)."
            ),
            (
                "Cumulative 7-Day Fatigue & Sleep Debt Forecaster",
                "Calculates cumulative sleep deficit from a 7-hour baseline to generate the Operational Strain Index (0-100%) and personalized recovery advisories."
            ),
            (
                "Proactive Burnout Prevention for Field Command",
                "Alerts unit officers when soldier fatigue crosses safety thresholds (above 65%), preventing weapon-handling mishaps and guard duty vigilance failures."
            )
        ]
    )

    # =========================================================================
    # SLIDE 7: CRISIS SOS & STABILIZATION
    # =========================================================================
    add_content_slide(
        7,
        "Emergency Response",
        "Life-Saving Crisis SOS & 4-7-8 Breathing Tool",
        [
            (
                "1-Click Defense Crisis SOS Emergency Beacon",
                "Provides instant one-tap dialing to Army Emergency Helpline (1902) and National Tele-MANAS (14416) during severe psychological breakdown."
            ),
            (
                "Command Dashboard Crimson Red Alert Intercept",
                "Triggering the SOS beacon immediately broadcasts high-priority pulsing red alert banners across officer and headquarters screens for rapid intervention."
            ),
            (
                "Tactical 4-7-8 Breathing & Panic De-Escalation Engine",
                "Interactive animated guide (4s Inhale, 7s Hold, 8s Exhale) activating the parasympathetic nervous system to quickly lower elevated heart rate and anxiety."
            )
        ]
    )

    # =========================================================================
    # SLIDE 8: OFFICER & COMMAND HQ SUITE
    # =========================================================================
    add_content_slide(
        8,
        "Command Suite",
        "Welfare Officer Supervision & Senior Admin Controls",
        [
            (
                "Live Personnel Monitoring Roster & Search Filters",
                "Enables Welfare Officers to view all squad members, filter by risk level, and immediately isolate soldiers requiring urgent clinical attention."
            ),
            (
                "Confidential Clinical Counseling Notes Engine",
                "Provides structured documentation for officer counseling sessions, tracking clinical observations, follow-up flags, and review schedules securely."
            ),
            (
                "Dynamic Command Master PIN & Senior HQ Analytics",
                "Senior Admins monitor army-wide risk distributions, review immutable activity audit trails, and rotate the 4-digit Master PIN in real-time."
            )
        ]
    )

    # =========================================================================
    # SLIDE 9: TECHNOLOGY STACK & SECURITY
    # =========================================================================
    add_content_slide(
        9,
        "Technology & Security",
        "Robust Tech Stack, Encryption & Vernacular Support",
        [
            (
                "Full-Stack Development Framework",
                "Engineered using Python Flask micro-framework, SQLite3 database, private Vercel Blob cloud storage, ReportLab PDF generators, and Vanilla CSS3/JS."
            ),
            (
                "Defense-Grade Security Architecture",
                "Enforces PBKDF2-HMAC-SHA256 password salting, parameterized SQL protection, self-service unit password recovery, and append-only activity logging."
            ),
            (
                "Zero-Latency Multi-Language Vernacular Switcher",
                "Client-side localization supporting English, Tamil (தமிழ்), and Hindi (हिंदी) stored in localStorage, making mental health care accessible to every jawan."
            )
        ]
    )

    # =========================================================================
    # SLIDE 10: CONCLUSION & FUTURE SCOPE
    # =========================================================================
    add_content_slide(
        10,
        "Conclusion",
        "Project Summary & Future Research Directions",
        [
            (
                "Production-Ready Deployment & Real-World Impact",
                "Successfully tested and deployed live on Vercel (https://apps-nu-henna.vercel.app), delivering a zero-latency digital welfare shield for defense forces."
            ),
            (
                "Friction-Free Direct Access & Complete Transparency",
                "Eliminated unnecessary authentication barriers for fast soldier reporting while maintaining strict command PIN gating and multi-role audit logging."
            ),
            (
                "Future Roadmap: Wearable IoT & Voice Sentiment AI",
                "Planned extensions include direct Bluetooth smartwatch hardware pairing, AI voice sentiment analysis during counseling, and edge-mesh offline database sync."
            )
        ]
    )

    prs.save(output_path)
    print(f"Presentation successfully created with {len(prs.slides)} slides at: {output_path}")

if __name__ == "__main__":
    out_file = sys.argv[1] if len(sys.argv) > 1 else "APPS_Mini_Project_Presentation.pptx"
    create_presentation(out_file)
